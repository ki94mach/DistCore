"""Background job runners for refresh and optimize."""

from __future__ import annotations

import io
from datetime import date
from typing import Any, Callable, Optional

from openpyxl import Workbook

from src.optimization.data import OptimizationSettings
from src.optimization.service import OptimizationService, RunRequest, RunResult
from src.orchestrator.pipelines.service import (
    PipelineService,
    RefreshAllRequest,
    RefreshAllResult,
)
from src.web.jobs.lock import SingleFlightLock
from src.web.jobs.store import JobStore


def serialize_refresh_result(result: RefreshAllResult) -> dict[str, Any]:
    return {
        "snapshot_date": result.snapshot_date.isoformat(),
        "succeeded": [item.value for item in result.succeeded],
        "failed": [item.value for item in result.failed],
        "is_successful": result.is_successful,
        "results": [
            {
                "pipeline": item.pipeline.value,
                "batch_type": item.batch_type,
                "snapshot_date": item.snapshot_date.isoformat(),
                "batch_id": item.batch_id,
                "status": item.status,
                "message": item.message,
            }
            for item in result.results
        ],
    }


def serialize_optimize_result(result: RunResult) -> dict[str, Any]:
    payload = result.to_dict()
    payload["shipments_full"] = [
        {
            "distributor_id": item.distributor_id,
            "distributor_name": item.distributor_name,
            "product_id": item.product_id,
            "product_name": item.product_name,
            "product_name_en": item.product_name_en,
            "quantity": item.quantity,
        }
        for item in result.shipments
    ]
    payload["table"] = {
        "columns": list(result.table.columns),
        "rows": [dict(row) for row in result.table.rows],
    }
    return payload


def build_excel_bytes(result: RunResult) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["status", result.status])
    summary.append(["solver_name", result.solver_name])
    summary.append(["objective", result.objective])
    summary.append(["is_optimal", result.is_optimal])
    summary.append(["is_feasible", result.is_feasible])
    summary.append(["total_shipments", result.summary.total_shipments])
    summary.append(["num_distributors", result.summary.num_distributors])
    summary.append(["num_products", result.summary.num_products])

    by_product = workbook.create_sheet("ByProduct")
    by_product.append(["product", "quantity"])
    for key, value in result.summary.shipments_by_product.items():
        by_product.append([key, value])

    by_distributor = workbook.create_sheet("ByDistributor")
    by_distributor.append(["distributor", "quantity"])
    for key, value in result.summary.shipments_by_distributor.items():
        by_distributor.append([key, value])

    table_sheet = workbook.create_sheet("Shipments")
    table_sheet.append(list(result.table.columns))
    for row in result.table.rows:
        table_sheet.append([row.get(column) for column in result.table.columns])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class JobRunner:
    """Execute refresh/optimize work while holding the single-flight lock."""

    def __init__(
        self,
        store: JobStore,
        lock: SingleFlightLock,
        *,
        pipeline_service_factory: Callable[[], PipelineService],
        optimization_service_factory: Callable[[], OptimizationService],
    ) -> None:
        self._store = store
        self._lock = lock
        self._pipeline_service_factory = pipeline_service_factory
        self._optimization_service_factory = optimization_service_factory

    def run_refresh(
        self,
        job_id: str,
        snapshot_date: date,
        include_deliveries: bool,
    ) -> None:
        try:
            self._store.mark_running(job_id, "Refreshing snapshots")
            service = self._pipeline_service_factory()
            result = service.refresh_all(
                RefreshAllRequest(
                    snapshot_date=snapshot_date,
                    include_deliveries=include_deliveries,
                    triggered_by="WEB_UI",
                )
            )
            self._store.write_result(job_id, serialize_refresh_result(result))
            if result.is_successful:
                self._store.mark_succeeded(job_id, "Refresh completed")
            else:
                failed = ", ".join(item.value for item in result.failed)
                self._store.mark_failed(job_id, f"Refresh partially failed: {failed}")
        except Exception as exc:
            self._store.mark_failed(job_id, str(exc))
        finally:
            self._lock.release(job_id)

    def run_optimize(
        self,
        job_id: str,
        snapshot_date: date,
        solver: str,
        settings: Optional[OptimizationSettings],
        settings_preset: Optional[str],
        solver_options: Optional[dict[str, dict[str, Any]]],
        include_export_variables: bool,
    ) -> None:
        try:
            self._store.mark_running(job_id, "Running optimization")
            service = self._optimization_service_factory()
            result = service.run(
                RunRequest(
                    snapshot_date=snapshot_date,
                    solver=solver,
                    settings=settings,
                    settings_preset=settings_preset,
                    solver_options=solver_options,
                    include_export_variables=include_export_variables,
                )
            )
            self._store.write_result(job_id, serialize_optimize_result(result))
            self._store.write_excel(job_id, build_excel_bytes(result))
            self._store.mark_succeeded(job_id, "Optimization completed")
        except Exception as exc:
            self._store.mark_failed(job_id, str(exc))
        finally:
            self._lock.release(job_id)
