"""Excel (.xlsx) export for optimization run results."""

from __future__ import annotations

from dataclasses import dataclass, fields
from datetime import date
from pathlib import Path
from typing import Any, Mapping, Optional, Sequence, Union

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.optimization.data import OptimizationSettings
from src.optimization.explain import (
    CONSTRAINT_COMPARISON_COLUMNS,
    CONSTRAINT_GROUP_LABELS,
    ROW_DETAIL_COLUMN_NAMES,
    SHIPMENTS_BASE_COLUMN_NAMES,
    constraint_group_for_column,
    shipment_column_label,
)
from src.optimization.service import RunResult

SHIPMENTS_BASE_COLUMNS = SHIPMENTS_BASE_COLUMN_NAMES
SHIPMENTS_DETAIL_COLUMNS = ROW_DETAIL_COLUMN_NAMES
SHIPMENTS_TABLE_COLUMNS = SHIPMENTS_BASE_COLUMNS + SHIPMENTS_DETAIL_COLUMNS
# Backward-compatible alias: detail columns excluding explanation
SHIPMENTS_VARIABLE_COLUMNS = SHIPMENTS_DETAIL_COLUMNS[1:]

_HEADER_FILL = PatternFill("solid", fgColor="102A43")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_GROUP_BAND_FONT = Font(bold=True, color="FFFFFF")
_GROUP_FILLS: dict[str, PatternFill] = {
    "result": PatternFill("solid", fgColor="102A43"),
    "explanation": PatternFill("solid", fgColor="334E68"),
    "inputs": PatternFill("solid", fgColor="486581"),
    "demand_coverage": PatternFill("solid", fgColor="2CB1BC"),
    "target_units": PatternFill("solid", fgColor="F4A261"),
    "delivery_history": PatternFill("solid", fgColor="627D98"),
    "delivery_smoothing": PatternFill("solid", fgColor="9B8BBA"),
    "factory_supply": PatternFill("solid", fgColor="829AB1"),
}
_COMPARISON_HEADER_FONT = Font(bold=True, color="102A43")
_COMPARISON_HEADER_FILLS: dict[str, PatternFill] = {
    "demand_coverage": PatternFill("solid", fgColor="B8ECF2"),
    "target_units": PatternFill("solid", fgColor="FDE4CF"),
}
_DEFAULT_COMPARISON_FILL = PatternFill("solid", fgColor="EEF3F7")
_ALT_ROW_FILL = PatternFill("solid", fgColor="F0F4F8")
_THIN_BORDER = Border(
    left=Side(style="thin", color="D9E2EC"),
    right=Side(style="thin", color="D9E2EC"),
    top=Side(style="thin", color="D9E2EC"),
    bottom=Side(style="thin", color="D9E2EC"),
)
_MIN_COL_WIDTH = 10
_MAX_COL_WIDTH = 42


@dataclass(frozen=True)
class ExportRequestMeta:
    """Request echo written to the Parameters sheet."""

    snapshot_date: date
    solver: str
    settings: Optional[OptimizationSettings] = None
    settings_preset: Optional[str] = None
    solver_options: Optional[Mapping[str, Mapping[str, Any]]] = None
    include_export_variables: bool = False


def _excel_value(value: Any) -> Any:
    """Convert Python values to Excel-safe cell contents (None -> blank)."""
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _append_kv(sheet: Worksheet, key: str, value: Any) -> None:
    sheet.append([key, _excel_value(value)])


def _display_width(value: Any) -> int:
    text = "" if value is None else str(value)
    # Rough Excel width: Persian/CJK chars count a bit wider.
    width = 0.0
    for char in text:
        width += 1.7 if ord(char) > 255 else 1.0
    return int(width) + 2


def _autosize_columns(
    sheet: Worksheet,
    *,
    min_width: int = _MIN_COL_WIDTH,
    max_width: int = _MAX_COL_WIDTH,
) -> None:
    widths: dict[int, int] = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            col_idx = cell.column
            widths[col_idx] = max(widths.get(col_idx, min_width), _display_width(cell.value))
    for col_idx, width in widths.items():
        sheet.column_dimensions[get_column_letter(col_idx)].width = min(
            max(width, min_width), max_width
        )


def _style_header_row(sheet: Worksheet, row_index: int = 1) -> None:
    for cell in sheet[row_index]:
        if cell.value is None and cell.column > 1:
            continue
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _THIN_BORDER
    sheet.row_dimensions[row_index].height = 22


def _style_table_body(
    sheet: Worksheet,
    *,
    start_row: int,
    end_row: int,
    start_col: int,
    end_col: int,
) -> None:
    for row_idx in range(start_row, end_row + 1):
        fill = _ALT_ROW_FILL if (row_idx - start_row) % 2 == 1 else None
        for col_idx in range(start_col, end_col + 1):
            cell = sheet.cell(row=row_idx, column=col_idx)
            cell.border = _THIN_BORDER
            cell.alignment = Alignment(vertical="center")
            if fill is not None:
                cell.fill = fill


def _format_sheet_as_table(
    sheet: Worksheet,
    *,
    header_row: int = 1,
    freeze: bool = True,
) -> None:
    if sheet.max_row < header_row or sheet.max_column < 1:
        return
    _style_header_row(sheet, header_row)
    if sheet.max_row > header_row:
        _style_table_body(
            sheet,
            start_row=header_row + 1,
            end_row=sheet.max_row,
            start_col=1,
            end_col=sheet.max_column,
        )
    _autosize_columns(sheet)
    if freeze:
        sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1).coordinate
    sheet.sheet_view.showGridLines = False


def _format_kv_sections(sheet: Worksheet, section_header_rows: Sequence[int]) -> None:
    """Style key/value sheets that mix summary blocks and small tables."""
    header_set = set(section_header_rows)
    for row in sheet.iter_rows(min_row=1, max_col=max(sheet.max_column, 2)):
        row_idx = row[0].row
        values = [cell.value for cell in row]
        if all(value is None for value in values):
            continue
        if row_idx in header_set or (
            isinstance(values[0], str)
            and values[0] in {"shipments_by_product", "shipments_by_distributor"}
        ):
            for cell in row:
                if cell.value is None:
                    continue
                cell.fill = _HEADER_FILL
                cell.font = _HEADER_FONT
                cell.border = _THIN_BORDER
            continue
        # Alternating fill for data / kv rows
        if row_idx % 2 == 0:
            for cell in row:
                if cell.value is None and cell.column > 2:
                    continue
                cell.fill = _ALT_ROW_FILL
        for cell in row:
            if cell.value is not None or cell.column <= 2:
                cell.border = _THIN_BORDER
        # Bold keys in column A for kv pairs
        if row[0].value is not None and row_idx not in header_set:
            row[0].font = Font(bold=True)
    _autosize_columns(sheet)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False


def _build_constraint_group_row(column_keys: Sequence[str]) -> tuple[list[str], list[tuple[int, int]]]:
    """Build band labels for row 1 and 1-based column merge ranges."""
    labels: list[str] = []
    merges: list[tuple[int, int]] = []
    index = 0
    while index < len(column_keys):
        group_id = constraint_group_for_column(column_keys[index])
        start = index + 1
        end = start
        while end <= len(column_keys) and constraint_group_for_column(column_keys[end - 1]) == group_id:
            end += 1
        band_label = CONSTRAINT_GROUP_LABELS.get(group_id, group_id)
        labels.append(band_label)
        for _ in range(start + 1, end):
            labels.append("")
        if end - start > 1:
            merges.append((start, end - 1))
        index = end - 1
    return labels, merges


def _style_shipments_group_row(
    sheet: Worksheet,
    column_keys: Sequence[str],
    *,
    row_index: int = 1,
) -> None:
    index = 0
    while index < len(column_keys):
        group_id = constraint_group_for_column(column_keys[index])
        start_col = index + 1
        end_index = index
        while end_index < len(column_keys) and constraint_group_for_column(column_keys[end_index]) == group_id:
            end_index += 1
        fill = _GROUP_FILLS.get(group_id, _HEADER_FILL)
        for col_idx in range(start_col, end_index + 1):
            cell = sheet.cell(row=row_index, column=col_idx)
            cell.fill = fill
            cell.font = _GROUP_BAND_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = _THIN_BORDER
        index = end_index
    sheet.row_dimensions[row_index].height = 20


def _style_shipments_column_header_row(
    sheet: Worksheet,
    column_keys: Sequence[str],
    *,
    row_index: int = 2,
) -> None:
    comparison_by_group = CONSTRAINT_COMPARISON_COLUMNS
    for col_idx, column_key in enumerate(column_keys, start=1):
        cell = sheet.cell(row=row_index, column=col_idx)
        if cell.value is None:
            continue
        group_id = constraint_group_for_column(column_key)
        comparisons = comparison_by_group.get(group_id, frozenset())
        if column_key in comparisons:
            cell.fill = _COMPARISON_HEADER_FILLS.get(group_id, _DEFAULT_COMPARISON_FILL)
            cell.font = _COMPARISON_HEADER_FONT
        else:
            cell.fill = _GROUP_FILLS.get(group_id, _HEADER_FILL)
            cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _THIN_BORDER
    sheet.row_dimensions[row_index].height = 28


def _write_shipments(sheet: Worksheet, result: RunResult) -> None:
    columns = list(result.table.columns)
    group_labels, merge_ranges = _build_constraint_group_row(columns)
    sheet.append(group_labels)
    sheet.append([shipment_column_label(column) for column in columns])
    for row in result.table.rows:
        sheet.append([_excel_value(row.get(column)) for column in columns])
    for start_col, end_col in merge_ranges:
        sheet.merge_cells(
            start_row=1,
            start_column=start_col,
            end_row=1,
            end_column=end_col,
        )
    if sheet.max_row < 2 or sheet.max_column < 1:
        return
    _style_shipments_group_row(sheet, columns, row_index=1)
    _style_shipments_column_header_row(sheet, columns, row_index=2)
    if sheet.max_row > 2:
        _style_table_body(
            sheet,
            start_row=3,
            end_row=sheet.max_row,
            start_col=1,
            end_col=sheet.max_column,
        )
    _autosize_columns(sheet)
    sheet.freeze_panes = sheet.cell(row=3, column=1).coordinate
    sheet.sheet_view.showGridLines = False


def _write_summary(sheet: Worksheet, result: RunResult) -> None:
    _append_kv(sheet, "status", result.status)
    _append_kv(sheet, "objective", result.objective)
    _append_kv(sheet, "solver_name", result.solver_name)
    _append_kv(sheet, "is_optimal", result.is_optimal)
    _append_kv(sheet, "is_feasible", result.is_feasible)
    _append_kv(sheet, "total_shipments", result.summary.total_shipments)
    _append_kv(sheet, "num_distributors", result.summary.num_distributors)
    _append_kv(sheet, "num_products", result.summary.num_products)

    sheet.append([])
    sheet.append(["shipments_by_product", "quantity"])
    product_header_row = sheet.max_row
    for product, quantity in result.summary.shipments_by_product.items():
        sheet.append([_excel_value(product), _excel_value(quantity)])

    sheet.append([])
    sheet.append(["shipments_by_distributor", "quantity"])
    distributor_header_row = sheet.max_row
    for distributor, quantity in result.summary.shipments_by_distributor.items():
        sheet.append([_excel_value(distributor), _excel_value(quantity)])

    _format_kv_sections(sheet, (product_header_row, distributor_header_row))


def _resolved_settings(request: ExportRequestMeta) -> OptimizationSettings:
    return request.settings if request.settings is not None else OptimizationSettings()


def _write_parameters(sheet: Worksheet, request: ExportRequestMeta) -> None:
    _append_kv(sheet, "snapshot_date", request.snapshot_date)
    _append_kv(sheet, "solver", request.solver)
    _append_kv(sheet, "settings_preset", request.settings_preset or "")
    _append_kv(sheet, "include_export_variables", request.include_export_variables)

    settings = _resolved_settings(request)
    for field in fields(OptimizationSettings):
        _append_kv(sheet, f"settings.{field.name}", getattr(settings, field.name))

    options = request.solver_options or {}
    for solver_name, option_map in options.items():
        for option_name, option_value in option_map.items():
            _append_kv(sheet, f"{solver_name}.{option_name}", option_value)

    # Treat first row as a pseudo-header band for key/value sheets
    if sheet.max_row >= 1:
        sheet.insert_rows(1)
        sheet.cell(row=1, column=1, value="parameter")
        sheet.cell(row=1, column=2, value="value")
    _format_sheet_as_table(sheet, header_row=1, freeze=True)


def build_optimization_workbook(
    result: RunResult,
    *,
    request: ExportRequestMeta,
) -> Workbook:
    """Build a workbook with Shipments, Summary, and Parameters sheets."""
    workbook = Workbook()
    shipments = workbook.active
    shipments.title = "Shipments"
    _write_shipments(shipments, result)

    summary = workbook.create_sheet("Summary")
    _write_summary(summary, result)

    parameters = workbook.create_sheet("Parameters")
    _write_parameters(parameters, request)
    return workbook


def write_optimization_xlsx(
    result: RunResult,
    path: Union[str, Path],
    *,
    request: ExportRequestMeta,
) -> Path:
    """Write the optimization workbook to ``path`` and return that path."""
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = build_optimization_workbook(result, request=request)
    workbook.save(output)
    return output


def download_filename(snapshot_date: date, solver: str) -> str:
    """Human-friendly attachment filename for HTTP downloads."""
    safe_solver = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in solver)
    return f"optimization_{snapshot_date.isoformat()}_{safe_solver}.xlsx"


__all__ = [
    "ExportRequestMeta",
    "SHIPMENTS_BASE_COLUMNS",
    "SHIPMENTS_DETAIL_COLUMNS",
    "SHIPMENTS_TABLE_COLUMNS",
    "SHIPMENTS_VARIABLE_COLUMNS",
    "build_optimization_workbook",
    "download_filename",
    "write_optimization_xlsx",
]
