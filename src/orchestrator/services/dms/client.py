"""DMS SharePoint client for listing and downloading files.

Auth modes:
  - Windows SSPI (domain-joined Windows, no password) — default on win32 when
    credentials are not provided
  - NTLM username/password — Linux servers and optional Windows fallback
  - Basic username/password — if the SharePoint/IIS site advertises Basic
"""

from __future__ import annotations

import io
import re
import sys
import warnings
from typing import Any, Dict, List, Optional
from urllib.parse import unquote, urljoin, urlparse

import pandas as pd
import requests

from src.orchestrator.services.dms.config import DmsConfigLoader

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
warnings.filterwarnings("ignore", message=".*Data Validation.*")
warnings.filterwarnings("ignore", message=".*Print area.*")

_EXCEL_HREF_PATTERN = re.compile(
    r'href="([^"]+\.(?:xlsx|xls))"',
    re.IGNORECASE,
)

_SUPPORTED_AUTH_MODES = frozenset({"auto", "sspi", "ntlm", "basic"})


def _get_sspi_auth():
    """Return Windows integrated auth handler (SSPI)."""
    if sys.platform != "win32":
        raise OSError(
            "DMS auth mode 'sspi' requires Windows integrated authentication. "
            "On Linux use NTLM/Basic with DISTCORE_DMS_USERNAME / "
            "DISTCORE_DMS_PASSWORD (or username/password in dms.yml)."
        )
    try:
        from requests_negotiate_sspi import HttpNegotiateAuth
    except ImportError as exc:
        raise ImportError(
            "requests-negotiate-sspi is required for DMS SSPI access. "
            "Install it with: pip install requests-negotiate-sspi"
        ) from exc
    return HttpNegotiateAuth()


def _get_ntlm_auth(username: str, password: str):
    try:
        from requests_ntlm import HttpNtlmAuth
    except ImportError as exc:
        raise ImportError(
            "requests-ntlm is required for DMS NTLM access. "
            "Install it with: pip install requests-ntlm"
        ) from exc
    return HttpNtlmAuth(username, password)


def _get_basic_auth(username: str, password: str):
    from requests.auth import HTTPBasicAuth

    return HTTPBasicAuth(username, password)


def resolve_dms_auth(
    *,
    username: Optional[str] = None,
    password: Optional[str] = None,
    auth_mode: Optional[str] = None,
    config_file_path=None,
):
    """
    Pick an auth handler for DMS requests.

    Explicit arguments override env/config. Mode ``auto`` chooses NTLM when
    username+password are present, otherwise SSPI on Windows.
    """
    settings = DmsConfigLoader.get_auth_settings(config_file_path)
    username = (username if username is not None else settings.get("username")) or None
    password = password if password is not None else settings.get("password")
    mode = (auth_mode or settings.get("auth") or "auto").strip().lower()

    if mode not in _SUPPORTED_AUTH_MODES:
        raise ValueError(
            f"Unsupported DMS auth mode '{mode}'. "
            f"Expected one of: {', '.join(sorted(_SUPPORTED_AUTH_MODES))}"
        )

    if mode == "auto":
        if username and password:
            mode = "ntlm"
        elif sys.platform == "win32":
            mode = "sspi"
        else:
            raise OSError(
                "DMS access on Linux requires NTLM or Basic credentials. "
                "Set DISTCORE_DMS_USERNAME and DISTCORE_DMS_PASSWORD, "
                "or username/password under the dms: key in dms.yml."
            )

    if mode == "sspi":
        return _get_sspi_auth()

    if not username or password is None:
        raise ValueError(
            f"DMS auth mode '{mode}' requires username and password "
            "(DISTCORE_DMS_USERNAME / DISTCORE_DMS_PASSWORD or dms.yml)."
        )

    if mode == "ntlm":
        return _get_ntlm_auth(username, password)
    return _get_basic_auth(username, password)


def _normalize_folder_url(folder_url: str) -> str:
    return folder_url.rstrip("/")


def _resolve_href(folder_url: str, href: str) -> str:
    if href.startswith("http://") or href.startswith("https://"):
        return href
    if href.startswith("/"):
        parsed = urlparse(folder_url)
        return f"{parsed.scheme}://{parsed.netloc}{href}"
    return urljoin(folder_url + "/", href)


def _filename_from_url(file_url: str) -> str:
    path = urlparse(file_url).path
    return unquote(path.rsplit("/", 1)[-1])


class DmsClient:
    """
    Client for listing and downloading files from DMS SharePoint folder pages.

    Folder pages are scraped for direct Excel file links. Authentication is
    SSPI on Windows by default, or NTLM/Basic when credentials are configured.
    """

    def __init__(
        self,
        auth=None,
        session: Optional[requests.Session] = None,
        *,
        username: Optional[str] = None,
        password: Optional[str] = None,
        auth_mode: Optional[str] = None,
        config_file_path=None,
    ):
        if auth is not None:
            self._auth = auth
        else:
            self._auth = resolve_dms_auth(
                username=username,
                password=password,
                auth_mode=auth_mode,
                config_file_path=config_file_path,
            )
        self._session = session or requests.Session()

    def _get(self, url: str, **kwargs) -> requests.Response:
        timeout = kwargs.pop("timeout", 60)
        allow_redirects = kwargs.pop("allow_redirects", True)
        response = self._session.get(
            url,
            auth=self._auth,
            timeout=timeout,
            allow_redirects=allow_redirects,
            **kwargs,
        )
        response.raise_for_status()
        return response

    def list_files(
        self,
        folder_url: str,
        pattern: Optional[str] = None,
    ) -> List[Dict[str, Any]]:
        """
        List Excel files in a DMS folder by parsing the SharePoint HTML page.

        Returns file metadata dicts with keys: name, path, size, _folder_url.
        """
        folder_url = _normalize_folder_url(folder_url)
        response = self._get(folder_url)

        files: List[Dict[str, Any]] = []
        seen_urls = set()

        for match in _EXCEL_HREF_PATTERN.finditer(response.text):
            file_url = _resolve_href(folder_url, match.group(1))
            if file_url in seen_urls:
                continue
            seen_urls.add(file_url)

            file_name = _filename_from_url(file_url)
            if pattern is not None and not re.search(pattern, file_name):
                continue

            files.append(
                {
                    "name": file_name,
                    "path": file_url,
                    "size": None,
                    "modified": None,
                    "_folder_url": folder_url,
                }
            )

        return files

    def download_file(
        self,
        file_path: str,
        folder_url: Optional[str] = None,
    ) -> bytes:
        """
        Download a file from DMS.

        file_path may be a full download URL (preferred) or a filename within
        folder_url (which triggers a folder listing lookup).
        """
        if file_path.startswith("http://") or file_path.startswith("https://"):
            download_url = file_path
        else:
            if not folder_url:
                raise ValueError(
                    "folder_url is required when file_path is not a full URL"
                )
            folder_url = _normalize_folder_url(folder_url)
            files = self.list_files(folder_url)
            matching = [
                file_info
                for file_info in files
                if file_info["name"] == file_path or file_info["path"].endswith(file_path)
            ]
            if not matching:
                available = [file_info["name"] for file_info in files[:10]]
                raise FileNotFoundError(
                    f"File '{file_path}' not found in DMS folder. "
                    f"Available files (first 10): {available}"
                )
            download_url = matching[0]["path"]

        response = self._get(download_url)
        return response.content

    def read_excel_table(
        self,
        file_path: str,
        sheet_name: str,
        table_name: str,
        folder_url: Optional[str] = None,
        file_name: Optional[str] = None,
    ) -> pd.DataFrame:
        """Read an Excel table (ListObject) from a DMS file into a DataFrame."""
        file_content = self.download_file(file_path, folder_url=folder_url)
        display_name = file_name or (
            _filename_from_url(file_path) if file_path.startswith("http") else file_path
        )

        try:
            import openpyxl

            workbook = openpyxl.load_workbook(
                io.BytesIO(file_content),
                read_only=False,
                data_only=True,
            )

            if sheet_name not in workbook.sheetnames:
                workbook.close()
                raise ValueError(
                    f"Worksheet named '{sheet_name}' not found in file '{display_name}'. "
                    f"Available sheets: {', '.join(workbook.sheetnames)}"
                )

            worksheet = workbook[sheet_name]

            if not worksheet.tables:
                workbook.close()
                raise ValueError(
                    f"No tables found in sheet '{sheet_name}' of file '{display_name}'."
                )

            if table_name not in worksheet.tables:
                workbook.close()
                available_tables = list(worksheet.tables.keys())
                raise ValueError(
                    f"Table named '{table_name}' not found in sheet '{sheet_name}' "
                    f"of file '{display_name}'. "
                    f"Available tables: {', '.join(available_tables) if available_tables else 'None'}"
                )

            table = worksheet.tables[table_name]
            table_range = table.ref
            if "!" in table_range:
                table_range = table_range.split("!")[1]

            start_cell, end_cell = table_range.split(":")
            cell_pattern = re.compile(r"^([A-Z]+)(\d+)$")
            start_match = cell_pattern.match(start_cell)
            end_match = cell_pattern.match(end_cell)

            if not start_match or not end_match:
                workbook.close()
                raise ValueError(
                    f"Invalid table range format '{table_range}' in file '{display_name}'. "
                    f"Expected format like 'A1:Z100'"
                )

            start_col_letter, start_row_str = start_match.groups()
            end_col_letter, end_row_str = end_match.groups()
            start_row = int(start_row_str)
            end_row = int(end_row_str)
            start_col = openpyxl.utils.column_index_from_string(start_col_letter) - 1
            end_col = openpyxl.utils.column_index_from_string(end_col_letter)

            headers = []
            for col_idx in range(start_col, end_col):
                cell = worksheet.cell(row=start_row, column=col_idx + 1)
                headers.append(
                    cell.value if cell.value is not None else f"Column{col_idx + 1}"
                )

            data = []
            for row_idx in range(start_row + 1, end_row + 1):
                row_data = []
                for col_idx in range(start_col, end_col):
                    cell = worksheet.cell(row=row_idx, column=col_idx + 1)
                    row_data.append(cell.value)
                data.append(row_data)

            workbook.close()
            df = pd.DataFrame(data, columns=headers)
            df = df.dropna(how="all")

            if df.empty:
                raise ValueError(
                    f"Table '{table_name}' in file '{display_name}' has no data rows"
                )

            return df

        except ImportError as exc:
            raise ImportError("openpyxl is required to read Excel tables") from exc
