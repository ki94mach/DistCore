"""DMS and Excel file processing utilities for distributor deliveries pipeline."""

import io
import re
import unicodedata
from typing import Callable, List, Optional, Tuple
import pandas as pd
import openpyxl

from src.orchestrator.pipelines.utils.excel_utils import parse_excel_date, safe_int
from src.orchestrator.services.dms import DmsClient


def concat_delivery_dataframes(dataframes: List[pd.DataFrame]) -> pd.DataFrame:
    """
    Concatenate per-file delivery DataFrames without pandas FutureWarnings.

    Drops all-NA columns per frame and aligns columns before concat so mixed
    Excel schemas across factory files do not trigger dtype inference warnings.
    """
    if not dataframes:
        raise ValueError("No dataframes to concatenate")

    if len(dataframes) == 1:
        return dataframes[0].copy()

    prepared: List[pd.DataFrame] = []
    all_columns: List[str] = []

    for dataframe in dataframes:
        frame = dataframe.dropna(axis=1, how='all')
        for column in frame.columns:
            if column not in all_columns:
                all_columns.append(column)
        prepared.append(frame)

    aligned = [frame.reindex(columns=all_columns) for frame in prepared]
    return pd.concat(aligned, ignore_index=True, sort=False)


def extract_factory_name_from_filename(filename: str) -> Optional[str]:
    """Extract factory name from filename pattern: ... - [Factory Name] - 1404.xlsx"""
    pattern = r'دیتابیس تحویل به پخش ها\s*-\s*([^-]+?)\s*-\s*\d+\.(xlsx|xls)$'
    match = re.search(pattern, filename)
    if match:
        return match.group(1).strip()
    return None


def find_matching_table(
    dms_client: DmsClient,
    file_info: dict,
    factory_name: str,
    folder_url: str,
) -> Optional[Tuple[str, str]]:
    """Find the Excel table (ListObject) matching the factory name."""
    file_content = dms_client.download_file(
        file_info['path'],
        folder_url=folder_url,
    )
    workbook = openpyxl.load_workbook(io.BytesIO(file_content), read_only=False)
    available_sheets = workbook.sheetnames

    factory_stripped = factory_name.strip()
    factory_cleaned = ''.join(
        char for char in factory_stripped if unicodedata.category(char) != 'Cf'
    )

    for sheet_name in available_sheets:
        worksheet = workbook[sheet_name]
        if not worksheet.tables:
            continue

        for table_name in worksheet.tables.keys():
            table_stripped = table_name.strip()
            table_cleaned = ''.join(
                char for char in table_stripped if unicodedata.category(char) != 'Cf'
            )

            if table_stripped == factory_stripped:
                workbook.close()
                return (sheet_name, table_name)

            if table_cleaned == factory_cleaned:
                workbook.close()
                return (sheet_name, table_name)

            table_no_spaces = table_cleaned.replace(' ', '').replace('\u200C', '').replace('\u200D', '')
            factory_no_spaces = factory_cleaned.replace(' ', '').replace('\u200C', '').replace('\u200D', '')
            if table_no_spaces == factory_no_spaces:
                workbook.close()
                return (sheet_name, table_name)

            for norm_form in ['NFC', 'NFD', 'NFKC', 'NFKD']:
                normalized_table = unicodedata.normalize(norm_form, table_stripped)
                normalized_factory = unicodedata.normalize(norm_form, factory_stripped)

                if normalized_table == normalized_factory:
                    workbook.close()
                    return (sheet_name, table_name)

                if normalized_factory in normalized_table or normalized_table in normalized_factory:
                    workbook.close()
                    return (sheet_name, table_name)

            if factory_cleaned in table_cleaned or table_cleaned in factory_cleaned:
                workbook.close()
                return (sheet_name, table_name)

            if table_cleaned.lower() == factory_cleaned.lower():
                workbook.close()
                return (sheet_name, table_name)

    workbook.close()
    return None


def normalize_dataframe_columns(df: pd.DataFrame, expected_columns: List[str]) -> pd.DataFrame:
    column_mapping = {}
    for col in df.columns:
        normalized = unicodedata.normalize('NFC', str(col).strip())
        column_mapping[col] = normalized

    df = df.rename(columns=column_mapping)

    for expected_col in expected_columns:
        if expected_col not in df.columns:
            for actual_col in df.columns:
                if (
                    unicodedata.normalize('NFC', actual_col.strip())
                    == unicodedata.normalize('NFC', expected_col.strip())
                ):
                    df = df.rename(columns={actual_col: expected_col})
                    break

    return df


def validate_dataframe_columns(df: pd.DataFrame, expected_columns: List[str]) -> None:
    missing_columns = [col for col in expected_columns if col not in df.columns]
    if missing_columns:
        raise ValueError(
            f"Missing {len(missing_columns)} expected column(s): {missing_columns}"
        )


def validate_dataframe_has_data(df: pd.DataFrame, data_columns: List[str]) -> None:
    if df.empty:
        raise ValueError("DataFrame is empty after loading Excel files.")

    if data_columns:
        non_null_count = df[data_columns].notna().any(axis=1).sum()
        if non_null_count == 0:
            raise ValueError("No data rows found in DataFrame. All values are null/empty.")


def filter_empty_rows(df: pd.DataFrame, data_columns: List[str]) -> pd.DataFrame:
    if not data_columns:
        return df

    has_data_mask = df[data_columns].notna().any(axis=1)
    for col in data_columns:
        has_data_mask = has_data_mask | (df[col].astype(str).str.strip() != '')

    return df[has_data_mask].copy()


def normalize_delivery_jalali_yyyymm(month, source_file: str = "") -> Optional[int]:
    """
    Normalize Excel ماه to Jalali YYYYMM.

    Accepts YYYYMM, YYYYMMDD, or bare month 1-12 (year taken from source_file name).
    """
    if month is None:
        return None
    try:
        value = int(month)
    except (TypeError, ValueError):
        return None

    if value >= 1000000:
        return value // 100
    if value >= 100000:
        return value
    if 1 <= value <= 12:
        match = re.search(r"(14\d{2})", source_file or "")
        if match:
            return int(match.group(1)) * 100 + value
        return None
    return None


def safe_get_column(row: pd.Series, column_name: str, default=None):
    if column_name not in row.index:
        return default
    value = row[column_name]
    if pd.isna(value):
        return default
    return value


def transform_distributor_delivery_row(
    row: pd.Series,
    persian_columns: List[str],
) -> tuple:
    request_date = parse_excel_date(safe_get_column(row, 'تاریخ درخواست'))
    expiry_date = parse_excel_date(safe_get_column(row, 'تاریخ انقضاء'))
    release_date = parse_excel_date(safe_get_column(row, 'تاریخ ریلیز'))
    delivery_date = parse_excel_date(safe_get_column(row, 'تاریخ تحویل'))

    delivered_qty = safe_int(safe_get_column(row, 'تعداد تحویلی'))
    delivered_qty_round_up = safe_int(safe_get_column(row, 'تعداد تحویلی رند بالا'))
    delivered_qty_round_down = safe_int(safe_get_column(row, 'تعداد تحویلی رند پایین'))
    warehouse_exit_qty = safe_int(safe_get_column(row, 'تعداد خروج از انبار'))
    routine_delivery_qty = safe_int(safe_get_column(row, 'تعداد تحویل روتین'))
    supply_chain_request_qty = safe_int(safe_get_column(row, 'تعداد درخواست زنجیره تامین'))
    days_between_delivery_release = safe_int(
        safe_get_column(row, 'تعداد روز بین تاریخ تحویلی و ریلیز')
    )

    source_file = safe_get_column(row, 'source_file', '') or ''
    month = normalize_delivery_jalali_yyyymm(
        safe_get_column(row, 'ماه'),
        source_file,
    )

    def safe_str(val, default=''):
        if val is None or pd.isna(val):
            return default
        return str(val)

    return (
        safe_str(safe_get_column(row, 'کد دارو')),
        safe_str(safe_get_column(row, 'کالا')),
        safe_str(safe_get_column(row, 'نام پخش')),
        safe_str(safe_get_column(row, 'شماره بچ')),
        expiry_date,
        delivered_qty,
        delivered_qty_round_up,
        delivered_qty_round_down,
        request_date,
        delivery_date,
        month,
        safe_str(safe_get_column(row, 'شماره نامه خروج از انبار')),
        warehouse_exit_qty,
        safe_str(safe_get_column(row, 'جزئیات خروج از انبار')),
        safe_str(safe_get_column(row, 'وضعیت رسید')),
        release_date,
        days_between_delivery_release,
        routine_delivery_qty,
        supply_chain_request_qty,
        safe_str(safe_get_column(row, 'کارخانه')),
        None,
        safe_str(safe_get_column(row, 'توضیحات')),
        source_file,
        None,
    )


def load_excel_files_from_dms(
    dms_client: DmsClient,
    folder_url: str,
    file_pattern: str,
    expected_files: List[str],
    ignored_files: List[str],
    persian_columns: List[str],
    header_row: int = 7,
    log_fn: Optional[Callable[[str], None]] = None,
    log_label: str = "files",
) -> pd.DataFrame:
    """
    Load and concatenate Excel files from a DMS folder.

    When log_fn is provided, emits a summary line: expected vs loaded file count.
    """
    del header_row  # table-based read ignores header row

    files = dms_client.list_files(folder_url=folder_url, pattern=file_pattern)
    if not files:
        raise ValueError(
            f"No files found in DMS folder matching pattern '{file_pattern}'"
        )

    files_to_process = [file_info for file_info in files if file_info['name'] not in (ignored_files or [])]
    if not files_to_process:
        raise ValueError("No files to process after filtering ignored files")

    dataframes = []
    skipped_files = []

    for file_info in files_to_process:
        file_name = file_info['name']
        try:
            factory_name = extract_factory_name_from_filename(file_name)
            if not factory_name:
                skipped_files.append(file_name)
                continue

            matching_table = find_matching_table(
                dms_client,
                file_info,
                factory_name,
                folder_url,
            )
            if not matching_table:
                skipped_files.append(file_name)
                continue

            sheet_name, table_name = matching_table
            df = dms_client.read_excel_table(
                file_info['path'],
                sheet_name=sheet_name,
                table_name=table_name,
                folder_url=folder_url,
                file_name=file_name,
            )

            if df.empty or len(df) == 0:
                skipped_files.append(file_name)
                continue

            df['source_file'] = file_name
            dataframes.append(df)

        except (ValueError, Exception):
            skipped_files.append(file_name)

    if not dataframes:
        raise ValueError(
            f"No valid Excel files could be loaded. "
            f"Total files found: {len(files)}, Skipped: {len(skipped_files)}"
        )

    non_empty_dataframes = [
        dataframe for dataframe in dataframes
        if not dataframe.empty and len(dataframe) > 0
    ]
    if not non_empty_dataframes:
        raise ValueError(
            f"No valid data found in Excel files. "
            f"Total files found: {len(files)}, Skipped: {len(skipped_files)}"
        )

    df = concat_delivery_dataframes(non_empty_dataframes)
    df = normalize_dataframe_columns(df, persian_columns)
    validate_dataframe_columns(df, persian_columns)
    validate_dataframe_has_data(df, persian_columns)
    df = filter_empty_rows(df, persian_columns)

    columns_to_keep = persian_columns.copy()
    if 'source_file' in df.columns:
        columns_to_keep.append('source_file')
    df = df[columns_to_keep]

    loaded_files = df['source_file'].unique().tolist() if 'source_file' in df.columns else []
    expected_count = len(expected_files) if expected_files else 0
    if log_fn is not None:
        log_fn(f"Expected {expected_count} {log_label}, loaded {len(loaded_files)} {log_label}")

    return df


def transform_dataframe_to_fact_rows(
    df: pd.DataFrame,
    persian_columns: List[str],
) -> List[tuple]:
    missing_cols = [col for col in persian_columns if col not in df.columns]
    if missing_cols:
        raise ValueError(f"Missing required columns in DataFrame: {missing_cols}")

    staging_rows = []
    total_rows = len(df)
    progress_interval = max(1000, total_rows // 10)

    for idx, (_, row) in enumerate(df.iterrows(), 1):
        if idx % progress_interval == 0 or idx == total_rows:
            print(
                f"  Transforming row {idx}/{total_rows} ({idx * 100 // total_rows}%)...",
                end='\r',
                flush=True,
            )

        fact_row = transform_distributor_delivery_row(row, persian_columns)
        staging_rows.append(fact_row)

    print(" " * 60, end='\r', flush=True)
    return staging_rows


transform_dataframe_to_staging_rows = transform_dataframe_to_fact_rows
